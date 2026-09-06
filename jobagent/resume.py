"""Turn your resume file into plain text once, then cache it."""

from __future__ import annotations

import json
from pathlib import Path

CACHE_NAME = ".resume_cache.json"


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix in (".docx", ".doc"):
        text = _from_docx(path)
    elif suffix in (".txt", ".md"):
        text = path.read_text()
    else:
        raise ValueError(f"unsupported resume format: {suffix}. Use pdf, docx, txt or md.")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    return text.strip()


def load(path: Path, cache_dir: Path) -> str:
    """Extract with a cache keyed on the file mtime and size."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_file = cache_dir / CACHE_NAME
    stat = path.stat()
    key = f"{path}:{stat.st_mtime_ns}:{stat.st_size}"
    if cache_file.exists():
        try:
            cached = json.loads(cache_file.read_text())
            if cached.get("key") == key:
                return cached["text"]
        except (ValueError, KeyError):
            pass
    text = extract(path)
    cache_file.write_text(json.dumps({"key": key, "text": text}))
    return text
